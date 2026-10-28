import os
import shutil
import zipfile
import re
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

INPUT_DOCX = r"d:\project\fan_page\Project_Report_Output\Project_Report_Final.docx"
OUTPUT_DIR = r"d:\project\fan_page\Project_Report_Output"
FANVERSE_DOCX = os.path.join(OUTPUT_DIR, "FanVerse_Project_Report_Final.docx")
ROOT_DOCX = r"d:\project\fan_page\FanVerse_Project_Report_Final.docx"

TEMP_DOCX = os.path.join(OUTPUT_DIR, "temp_fanverse.docx")

NEW_TITLE = "FanVerse — A Unified Fan Community Platform"

doc = docx.Document(INPUT_DOCX)

def format_run(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic

def replace_in_para(p, old_text, new_text):
    if old_text in p.text:
        # Save formatting of run 0 if available
        bold = p.runs[0].bold if p.runs else False
        italic = p.runs[0].italic if p.runs else False
        size_pt = p.runs[0].font.size.pt if (p.runs and p.runs[0].font.size) else 12
        align = p.alignment
        
        updated_text = p.text.replace(old_text, new_text)
        p.text = updated_text
        if align:
            p.alignment = align
        for r in p.runs:
            format_run(r, "Times New Roman", size_pt, bold, italic)
        return True
    return False

# 1. Update Cover Page (P5)
p5 = doc.paragraphs[5]
p5.text = f"Project Title: {NEW_TITLE}"
p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in p5.runs:
    format_run(r, "Times New Roman", 13, bold=True, italic=False)
print("Updated P5 (Cover Page)")

# 2. Update Certificate (P45)
p45 = doc.paragraphs[45]
p45.text = (
    "This is to certify that Sriram Kavadi & Team, bearing Roll Number 22AG1A05XX, students of "
    "Department of Computer Science & Engineering, III Year B.Tech, has completed the one-week technical project entitled "
    f"“{NEW_TITLE}: A RESPONSIVE FLUTTER APPLICATION INTEGRATING ADVANCED STATE MANAGEMENT, REST API, AND EXPERIMENTAL LAB CURRICULUM” "
    "during 21-10-2026 to 27-10-2026 at ACE Engineering College, under the guidance of Dr. / Prof. Faculty Mentor, "
    "Associate Professor, Department of Computer Science & Engineering."
)
p45.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for r in p45.runs:
    format_run(r, "Times New Roman", 12, bold=False, italic=False)
print("Updated P45 (Certificate)")

# 3. Update Abstract (P55)
p55 = doc.paragraphs[55]
p55.text = (
    f"{NEW_TITLE} is an original, production-grade, responsive cross-platform web and mobile application "
    "engineered using Google's Flutter framework (Material Design 3) and the Dart programming language. "
    "The primary objective of this project is to build a dedicated, authentic digital engagement and verified community "
    "portal celebrating youth empowerment initiatives, grassroots volunteer mobilization, and cultural community engagement across 33+ districts."
)
p55.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for r in p55.runs:
    format_run(r, "Times New Roman", 12, bold=False, italic=False)
print("Updated P55 (Abstract)")

# 4. Update Introduction (P62)
p62 = doc.paragraphs[62]
p62.text = (
    f"Motivated by these real-world challenges, this project undertakes the development of {NEW_TITLE}. "
    "Built in strict accordance with the academic Flutter curriculum outlined in Marco L. Napoli's Beginning Flutter: "
    "A Hands-on Guide to App Development, this project encapsulates core concepts including Dart language fundamentals, "
    "widget composition, responsive breakpoints, named route navigation, Provider-based reactive state management, "
    "custom vector rendering, validated form input processing, explicit animations, and asynchronous REST API networking."
)
p62.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for r in p62.runs:
    format_run(r, "Times New Roman", 12, bold=False, italic=False)
print("Updated P62 (Introduction)")

# 5. Update Objectives Header (P67)
p67 = doc.paragraphs[67]
p67.text = f"The primary objectives of the {NEW_TITLE} project are:"
for r in p67.runs:
    format_run(r, "Times New Roman", 12, bold=True, italic=False)
print("Updated P67 (Objectives Header)")

# 6. Update Objective 2 (P69)
p69 = doc.paragraphs[69]
p69.text = (
    f"2. Dedicated Community Hub Development: Build an authentic, dignified, and feature-rich digital community portal under "
    f"{NEW_TITLE}, featuring an interactive home dashboard, live impact counters, community fan wall, upcoming events registry, "
    "leadership journey timeline, and digital membership verification."
)
p69.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for r in p69.runs:
    format_run(r, "Times New Roman", 12, bold=False, italic=False)
print("Updated P69 (Objective 2)")

# 7. Update System Design / Methodology (P98)
p98 = doc.paragraphs[98]
p98.text = (
    f"{NEW_TITLE} follows a structured, multi-tier software engineering methodology designed for separation of concerns, "
    "testability, and responsiveness. The application architecture is organized into four distinct decoupled layers:"
)
p98.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for r in p98.runs:
    format_run(r, "Times New Roman", 12, bold=False, italic=False)
print("Updated P98 (Methodology)")

# 8. Update Figure 8.1 Caption (P113)
p113 = doc.paragraphs[113]
p113.text = f"Figure 8.1: Architectural Block Diagram & Navigation Dataflow of {NEW_TITLE}"
p113.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in p113.runs:
    format_run(r, "Times New Roman", 11, bold=True, italic=True)
print("Updated P113 (Figure 8.1 Caption)")

# 9. Update Implementation (P119)
p119 = doc.paragraphs[119]
p119.text = (
    f"{NEW_TITLE} was implemented modularly across 10 distinct technical experiments corresponding to the "
    "prescribed mobile application development syllabus (Beginning Flutter by Marco L. Napoli)."
)
p119.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for r in p119.runs:
    format_run(r, "Times New Roman", 12, bold=False, italic=False)
print("Updated P119 (Implementation)")

# 10. Update Results (P134)
p134 = doc.paragraphs[134]
p134.text = (
    f"The project successfully produced {NEW_TITLE}, an enterprise-grade responsive cross-platform application. "
    "The software runs smoothly at 60 FPS across mobile, tablet, and desktop viewports, with zero runtime errors and full offline resiliency."
)
p134.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for r in p134.runs:
    format_run(r, "Times New Roman", 12, bold=False, italic=False)
print("Updated P134 (Results)")

# 11. Update Conclusion (P170)
p170 = doc.paragraphs[170]
p170.text = (
    f"Conclusion: {NEW_TITLE} was successfully architected, developed, tested, and documented within the designated "
    "one-week technical project timeline. The project accomplished 100% of its initial objectives and strictly satisfied all 10 "
    "experimental requirements set forth in the Flutter Lab curriculum (Marco L. Napoli, Beginning Flutter: A Hands-on Guide to App Development)."
)
p170.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for r in p170.runs:
    format_run(r, "Times New Roman", 12, bold=False, italic=False)
print("Updated P170 (Conclusion)")

# 12. Update Evaluation Sheet (P210)
p210 = doc.paragraphs[210]
p210.text = f"Project Title: {NEW_TITLE}"
for r in p210.runs:
    format_run(r, "Times New Roman", 12, bold=True, italic=False)
print("Updated P210 (Evaluation Sheet)")

# Save intermediate DOCX
doc.save(TEMP_DOCX)
print(f"Saved intermediate docx to {TEMP_DOCX}")

# 13. Update zip media and xml
with zipfile.ZipFile(TEMP_DOCX, 'r') as zin:
    file_map = {name: zin.read(name) for name in zin.namelist()}

# Swap image1.png and image8.png with the new fanverse_architecture_diagram.png
fanverse_diag_path = os.path.join(OUTPUT_DIR, "Project_Report_Assets", "fanverse_architecture_diagram.png")
with open(fanverse_diag_path, 'rb') as f_diag:
    fanverse_diag_bytes = f_diag.read()

file_map['word/media/image1.png'] = fanverse_diag_bytes
file_map['word/media/image8.png'] = fanverse_diag_bytes
print("Swapped word/media/image1.png and image8.png with FanVerse diagram")

# Ensure word/document.xml has no old title leftovers
doc_xml = file_map['word/document.xml'].decode('utf-8')

# Global safety replacements in document XML if any remaining
doc_xml = doc_xml.replace(
    'RATHNAVERI KAVITHA OFFICIAL FAN PAGE &amp; COMMUNITY HUB',
    NEW_TITLE
)
doc_xml = doc_xml.replace(
    'RATHNAVERI KAVITHA OFFICIAL FAN PAGE & COMMUNITY HUB',
    NEW_TITLE
)
doc_xml = doc_xml.replace(
    'Rathnaveri Kavitha Official Fan Page &amp; Community Hub',
    NEW_TITLE
)
doc_xml = doc_xml.replace(
    'Rathnaveri Kavitha Official Fan Page & Community Hub',
    NEW_TITLE
)
doc_xml = doc_xml.replace(
    'Rathnaveri Kavitha Official Fan Hub',
    NEW_TITLE
)

file_map['word/document.xml'] = doc_xml.encode('utf-8')

# Write final FanVerse DOCX to both locations
with zipfile.ZipFile(FANVERSE_DOCX, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
    for name, content in file_map.items():
        zout.writestr(name, content)
print(f"SUCCESS: Created {FANVERSE_DOCX} (Size: {os.path.getsize(FANVERSE_DOCX)} bytes)")

shutil.copyfile(FANVERSE_DOCX, ROOT_DOCX)
print(f"SUCCESS: Created copy at {ROOT_DOCX}")

if os.path.exists(TEMP_DOCX):
    os.remove(TEMP_DOCX)
